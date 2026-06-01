from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl.utils import get_column_letter
from datetime import datetime
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import pandas as pd
from src.ui.components.ui_style_table import columa_style
import locale


class IngresosDeptosDetalladosUMAReport:
    def __init__(self, datos: pd.DataFrame, municipio, titulo_reporte, fecha_ini, fecha_fin):
        self.datos = datos
        self.municipio = municipio
        self.fecha_ini = fecha_ini
        self.fecha_fin = fecha_fin
        self.titulo = titulo_reporte

    def generar_excel(self, ruta_salida="ingresos_deptos_detallado.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Ingresos_Depto_Det"

        # ===============================
        # ⭐ ESTILOS
        # ===============================
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        title_font = Font(size=16, bold=True)
        date_font = Font(size=12, italic=True)
        total_font = Font(size=11, italic=True, bold=True)

        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        align_center = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # ===============================
        # ⭐ TITULO Y FECHA
        # ===============================
        titulo = f"{self.municipio['NombreMuni']} - {self.municipio['NombreDepto']}"
        titulo_rpt = f"{self.titulo}"
        fecha = f"Fecha de elaboración: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        periodo = f"Periodo desde el {self.fecha_ini} al {self.fecha_fin}"

        ws.merge_cells("A1:L1")
        ws["A1"] = titulo
        ws["A1"].alignment = align_center
        ws["A1"].font = title_font

        ws.merge_cells("A2:L4")
        ws["A2"] = titulo_rpt
        ws["A2"].alignment = align_center
        ws["A2"].font = title_font

        ws.merge_cells("A5:L5")
        ws["A5"] = fecha
        ws["A5"].alignment = align_center
        ws["A5"].font = date_font

        ws.merge_cells("A6:L6")
        ws["A6"] = periodo
        ws["A6"].alignment = align_center
        ws["A6"].font = date_font

        fila_actual = 9  # Aquí irán los encabezados
        ws.freeze_panes = "A10"
        # ===============================
        # ⭐ ENCABEZADOS
        # ===============================
        columnas = ["No. RECIBO", "DNI", "NOMBRE DEL CONTRIBUYENTE", "LICENCIAS",
                    "BOSQUES Y DERIVADOS", "TASA AMBIENTAL", "MATRICULA MOTO SIERRA", "PERFORACION DE POZOS",
                    "CONSTNCIAS", "INSPECCIONES", "OTROS", "TOTAL"]

        col_value = 0
        for col in range(1, len(columnas) + 1):
            c = ws.cell(row=fila_actual, column=col, value=columnas[col_value])
            c.fill = header_fill
            c.font = header_font
            c.alignment = align_center
            c.border = border
            col_value += 1

        fila_actual += 1

        # ===============================
        # ⭐ LLENADO DE DATOS
        # ===============================
        total = 0.0
        total_otros = 0.0
        total_liciecias = 0.0
        total_bosques = 0.0
        total_tasa_ambiental = 0.0
        total_matricula_sierra = 0.0
        total_perforacion_pozo = 0.0
        total_constancia = 0.0
        total_liciencias = 0.0

        for index, item in self.datos.iterrows():
            total += item["TotalReciboPagado"]
            nombre_completo = f'{item["Pnombre"]} {item["SNombre"]} {item["PApellido"]} {item["SApellido"]}'
            num = item["NumRecibo"]
            fila_actual += 1
            total_sum = item["LicenciasExtranccion"] + item["BosquesDerivados"] + item["TasaAmbiental"] + \
                item["MatriculaMotoSierra"]+item["PerforacionPozos"] + \
                item["Constancias"]+item["Inspeccion"]
            otros_sum = 0.0
            if total_sum == item["TotalReciboPagado"]:
                if total_sum < item["TotalReciboPagado"]:
                    otros_sum = item["TotalReciboPagado"] - total_sum
                else:
                    otros_sum = 0
            total_otros = +otros_sum
            total_liciecias += item["LicenciasExtranccion"]
            total_bosques += item["BosquesDerivados"]
            total_tasa_ambiental += item["TasaAmbiental"]
            total_matricula_sierra += item["MatriculaMotoSierra"]
            total_perforacion_pozo += item["PerforacionPozos"]
            total_constancia += item["Constancias"]
            total_liciencias += item["Inspeccion"]

            ws.append([
                num,
                item["DNI"].strip(),
                nombre_completo,
                f'{item["LicenciasExtranccion"]:,.2f}',
                f'{item["BosquesDerivados"]:,.2f}',
                f'{item["TasaAmbiental"]:,.2f}',
                f'{item["MatriculaMotoSierra"]:,.2f}',
                f'{item["PerforacionPozos"]:,.2f}',
                f'{item["Constancias"]:,.2f}',
                f'{item["Inspeccion"]:,.2f}',
                f'{otros_sum:,.2f}',
                f'{item["TotalReciboPagado"]:,.2f}',

            ])

        ws.append([
            "",
            "",
            "Total Ingresos",
            f'{total_liciecias:,.2f}',
            f'{total_bosques:,.2f}',
            f'{total_tasa_ambiental:,.2f}',
            f'{total_matricula_sierra:,.2f}',
            f'{total_perforacion_pozo:,.2f}',
            f'{total_constancia:,.2f}',
            f'{total_liciencias:,.2f}',
            f'{total_otros:,.2f}',
            f'{total:,.2f}',
        ])

        ws.cell(row=fila_actual, column=2).font = total_font
        ws.cell(row=fila_actual, column=3).font = total_font
        ws.cell(row=fila_actual, column=4).font = total_font
        ws.cell(row=fila_actual, column=5).font = total_font

        # ===============================
        # ⭐ AJUSTE AUTOMÁTICO DE COLUMNAS
        # ===============================
        for col in ws.columns:
            max_length = 0
            column = col[0].column  # número de columna
            column_letter = get_column_letter(column)

            for cell in col[8:]:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = max_length + 2  # un pequeño margen
            ws.column_dimensions[column_letter].width = adjusted_width

        # ===============================
        # ⭐ GUARDAR
        # ===============================
        wb.save(ruta_salida)
        return ruta_salida
