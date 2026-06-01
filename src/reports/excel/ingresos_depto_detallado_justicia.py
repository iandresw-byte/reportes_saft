import pandas as pd
from openpyxl.utils import get_column_letter
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from src.reports.utils.manejador_data_frame import sumar_ingresos_departo_justicia


class IngresosDeptosDetalladosJusticiaReport:
    def __init__(self, datos: pd.DataFrame, municipio, titulo_reporte, fecha_ini, fecha_fin):
        self.datos = datos
        self.municipio = municipio
        self.fecha_ini = fecha_ini
        self.fecha_fin = fecha_fin
        self.titulo = titulo_reporte

    def generar_excel(self, ruta_salida="ingresos_deptos_detallado_justicia.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Ingresos_Depto_Det"

        # ===============================
        # ⭐ ESTILOS
        # ===============================
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        title_font = Font(size=16, bold=True)
        date_font = Font(size=12, italic=True)
        total_font = Font(size=11, italic=True, bold=True)

        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        align_center = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # ===============================
        # ⭐ TITULO Y FECHA
        # ===============================
        titulo = f"{self.municipio['NombreMuni']} - {self.municipio['NombreDepto']}"
        titulo_rpt = f"{self.titulo}"
        fecha = f"Fecha de elaboración: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        periodo = f"Periodo desde el {self.fecha_ini} al {self.fecha_fin}"

        ws.merge_cells("A1:O1")
        ws["A1"] = titulo
        ws["A1"].alignment = align_center
        ws["A1"].font = title_font

        ws.merge_cells("A2:O4")
        ws["A2"] = titulo_rpt
        ws["A2"].alignment = align_center
        ws["A2"].font = title_font

        ws.merge_cells("A5:O5")
        ws["A5"] = fecha
        ws["A5"].alignment = align_center
        ws["A5"].font = date_font

        ws.merge_cells("A6:O6")
        ws["A6"] = periodo
        ws["A6"].alignment = align_center
        ws["A6"].font = date_font

        fila_actual = 9  # Aquí irán los encabezados
        ws.freeze_panes = "A10"
        # ===============================
        # ⭐ ENCABEZADOS
        # ===============================
        columnas = ["No.", "No. RECIBO",
                    "DNI",
                    "NOMBRE DEL CONTRIBUYENTE",
                    "AUTORIZACIONES Y VISTOS BUENOS",
                    "CARTAS DE VENTA",
                    "GUIAS PARA TRANSPORTAR GANADO",
                    "PERMISOS PARA FORJAR MARCAS DE HERRAR",
                    "MATRÍCULA DE MARCAS DE HERRAR",
                    "MATRÍCULA ARMAS",
                    "LICENCIA BUHUNEROS",
                    "INSCRIPCIONES",
                    "CONSTANCIAS",
                    "OTROS",
                    "TOTAL"]

        col_value = 0
        for col in range(1, len(columnas) + 1):
            c = ws.cell(row=fila_actual, column=col, value=columnas[col_value])
            c.fill = header_fill
            c.font = header_font
            c.alignment = align_center
            c.border = border
            col_value += 1

        fila_actual += 1

        # ===============================
        # ⭐ LLENADO DE DATOS
        # ===============================

        datos = sumar_ingresos_departo_justicia(self.datos)
        for index, item in datos.iterrows():
            ixd = index+1
            nombre_completo = item["Nombre Completo"]
            num = item["No. Recibo"]
            fila_actual += 1

            ws.append([
                ixd,
                num,
                item["DNI"].strip(),
                nombre_completo,
                f'{item["Autorizaciones y Vistos Buenos"]:,.2f}',
                f'{item["Cartas de Venta"]:,.2f}',
                f'{item["Guias para Transportar Ganado"]:,.2f}',
                f'{item["Permiso Para Forjar Marcas de Herrar"]:,.2f}',
                f'{item["Matrícula de Marcas de Herrar"]:,.2f}',
                f'{item["Matrícula de Armas"]:,.2f}',
                f'{item["Licencia Buhuneros"]:,.2f}',
                f'{item["Inscripciones"]:,.2f}',
                f'{item["Constancias"]:,.2f}',
                f'{item["Otros"]:,.2f}',
                f'{item["Total Recibo"]:,.2f}',
            ])

        ws.cell(row=fila_actual, column=2).font = total_font
        ws.cell(row=fila_actual, column=3).font = total_font
        ws.cell(row=fila_actual, column=4).font = total_font
        ws.cell(row=fila_actual, column=5).font = total_font

        # ===============================
        # ⭐ AJUSTE AUTOMÁTICO DE COLUMNAS
        # ===============================
        for col in ws.columns:
            max_length = 0
            column = col[0].column  # número de columna
            column_letter = get_column_letter(column)

            for cell in col[8:]:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = max_length + 2  # un pequeño margen
            ws.column_dimensions[column_letter].width = adjusted_width

        # ===============================
        # ⭐ GUARDAR
        # ===============================
        wb.save(ruta_salida)
        return ruta_salida
