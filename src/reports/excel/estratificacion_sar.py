import pandas as pd
from openpyxl.utils import get_column_letter
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class EstratificacionSARReport:
    def __init__(self, datos: pd.DataFrame, municipio, titulo_reporte):
        self.datos = datos
        self.municipio = municipio
        self.titulo = titulo_reporte

    def generar_excel(self, ruta_salida="estratificacion_sar.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Estratificacion SAR"

        # ===============================
        # ⭐ ESTILOS
        # ===============================
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        title_font = Font(size=16, bold=True)
        date_font = Font(size=12, italic=True)

        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # ===============================
        # ⭐ TITULO Y FECHA
        # ===============================
        titulo = f"{self.municipio['NombreMuni']} - {self.municipio['NombreDepto']}"
        titulo_rpt = f"REPORTE DE - {self.titulo}"
        fecha = f"Fecha de elaboración: {datetime.now().strftime('%d/%m/%Y %H:%M')}"

        ws.merge_cells("A1:K1")
        ws["A1"] = titulo
        ws["A1"].alignment = align_center
        ws["A1"].font = title_font

        ws.merge_cells("A2:K2")
        ws["A2"] = titulo_rpt
        ws["A2"].alignment = align_center
        ws["A2"].font = title_font

        ws.merge_cells("A3:K3")
        ws["A3"] = fecha
        ws["A3"].alignment = align_center
        ws["A3"].font = date_font

        fila_actual = 5  # Aquí irán los encabezados
        ws.freeze_panes = "A6"
        # ===============================
        # ⭐ ENCABEZADOS
        # ===============================
        columnas = [
            "RTN", "RTM", "NOMBRE COMERCIAL", "NO PERMISO", "PERIODO ULT. PERMISO",
            "ACTIVIDAD ECONOMICA", "DIRECCION", "TELEFONO", "CONSTITUCION",
            "FECHA INICIO", "VALOR DECLARADO", "PERIODO DECLARADO", "ESTRATIFICACION"
        ]

        col_value = 0
        for col in range(1, len(columnas) + 1):
            c = ws.cell(row=fila_actual, column=col, value=columnas[col_value])
            c.fill = header_fill
            c.font = header_font
            c.alignment = align_center
            c.border = border
            col_value += 1

        fila_actual += 1

        # ===============================
        # ⭐ LLENADO DE DATOS
        # ===============================

        for index, item in self.datos.iterrows():

            if item["Constitucion"] == 0:
                tipo_persona = "S.A"
            elif item["Constitucion"] == 1:
                tipo_persona = "Comerciante Individual"
            elif item["Constitucion"] == 2:
                tipo_persona = "S. de R. L."
            elif item["Constitucion"] == 3:
                tipo_persona = "Grupo o Socios"
            else:
                tipo_persona = "No especificado"

            if item["val_declarado"] <= 1000000:
                estrato = "Microempresa"
            elif item["val_declarado"] >= 1000000.01 and item["val_declarado"] < 10000000:
                estrato = "Pequeña Empresa"
            elif item["val_declarado"] >= 10000000.01 and item["val_declarado"] < 50000000:
                estrato = "Mediana Empresa"
            elif item["val_declarado"] >= 50000000.01:
                estrato = "Empresa Grande"
            else:
                estrato = "No especificado"
            rtn = item["RTN"].strip() if item["RTN"] else ""
            ws.append([
                rtn,
                item["rtm"].strip(),
                item["nombreComercial"].strip(),
                "",
                item["ultPeriodoUltPermiso"],
                item["Actividad"].strip(),
                item["Direccion"].strip(),
                item["Telefono"].strip(),
                tipo_persona,
                item["FechaInicio"].strftime(
                    "%d/%m/%Y") if isinstance(item["FechaInicio"], pd.Timestamp) else item["FechaInicio"],

                f'{item['val_declarado']:,.2f}',
                item['utlAnioDeclarado'],
                estrato
            ])

        # ===============================
        # ⭐ AJUSTE AUTOMÁTICO DE COLUMNAS
        # ===============================
        for col in ws.columns:
            max_length = 0
            column = col[0].column  # número de columna
            column_letter = get_column_letter(column)

            for cell in col[5:]:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = max_length + 2  # un pequeño margen
            ws.column_dimensions[column_letter].width = adjusted_width

        # ===============================
        # ⭐ GUARDAR
        # ===============================
        wb.save(ruta_salida)
        return ruta_salida
