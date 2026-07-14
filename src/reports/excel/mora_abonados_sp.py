import pandas as pd
from openpyxl.utils import get_column_letter
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from src.reports.tablas.abonados_por_servicio import tabla_abonados_servicio_excel

from src.reports.utils.pintar_encabezado_excel import pintar_encabezado

from src.reports.utils.manejador_abonados_sp import sumar_mora_abonados_servicios


class AbonadosServiciosExcelReport:
    def __init__(self, datos: pd.DataFrame, municipio, titulo_reporte, fecha_ini=None, fecha_fin=None, depto=None, tipo_impuesto=None):
        self.datos = datos
        self.municipio = municipio
        self.fecha_ini = fecha_ini
        self.fecha_fin = fecha_fin
        self.titulo = titulo_reporte
        self.fila_num = 0
        self.tipo_impuesto = tipo_impuesto
        self.departamento = depto

    def generar_excel(self, ruta_salida="mora_abonados_por_servicio.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Mora_SP"

        # ===============================
        # ⭐ ESTILOS
        # ===============================

        title_font = Font(size=16, bold=True)
        date_font = Font(size=12, italic=True)
        total_font = Font(size=11, italic=True, bold=True)
        align_center = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

        # ===============================
        # ⭐ TITULO Y FECHA
        # ===============================
        titulo = f"{self.municipio['NombreMuni']} - {self.municipio['NombreDepto']}"
        titulo_rpt = f"{self.titulo}"
        fecha = f"Fecha de elaboración: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        periodo = None
        if self.fecha_ini:
            periodo = f"Periodo desde el {self.fecha_ini} al {self.fecha_fin}"
        
        

        if self.tipo_impuesto == '0':
            letra_col = 'K'
        elif self.tipo_impuesto == '1':
            letra_col = 'M'
        elif self.tipo_impuesto == '2,3':
            letra_col = 'N'
        elif self.tipo_impuesto == '4':
            letra_col = 'M'
        else:
            letra_col = 'O'

        ws.merge_cells(f"A1:{letra_col}1")
        ws["A1"] = titulo
        ws["A1"].alignment = align_center
        ws["A1"].font = title_font

        ws.merge_cells(f"A2:{letra_col}4")
        ws["A2"] = titulo_rpt
        ws["A2"].alignment = align_center
        ws["A2"].font = title_font

        ws.merge_cells(f"A5:{letra_col}5")
        ws["A5"] = fecha
        ws["A5"].alignment = align_center
        ws["A5"].font = date_font
        if periodo:
            ws.merge_cells(f"A6:{letra_col}6")
            ws["A6"] = periodo
            ws["A6"].alignment = align_center
            ws["A6"].font = date_font

        self.fila_num = 9  # Aquí irán los encabezados
        ws.freeze_panes = "A10"

        # ===============================
        # ⭐ LLENADO DE DATOS
        # ===============================
        columnas = []
        columnas.append("No.")
      
        datos = sumar_mora_abonados_servicios(
            self.datos)
        ws = pintar_encabezado(self, datos, columnas, ws)
        self.fila_num += 1
        ws = tabla_abonados_servicio_excel(
            self, datos, ws)
    
     

        ws.cell(row=self.fila_num, column=2).font = total_font
        ws.cell(row=self.fila_num, column=3).font = total_font
        ws.cell(row=self.fila_num, column=4).font = total_font
        ws.cell(row=self.fila_num, column=5).font = total_font

        # ===============================
        # ⭐ AJUSTE AUTOMÁTICO DE COLUMNAS
        # ===============================
        for col in ws.columns:
            max_length = 0
            column = col[0].column  # número de columna
            column_letter = get_column_letter(column)

            for cell in col[8:]:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = max_length + 2  # un pequeño margen
            ws.column_dimensions[column_letter].width = adjusted_width

        # ===============================
        # ⭐ GUARDAR
        # ===============================
        wb.save(ruta_salida)
        return ruta_salida
