import pandas as pd
from openpyxl.utils import get_column_letter
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from src.reports.utils.pintar_encabezado_excel import pintar_encabezado
from src.reports.utils.manejador_data_frame import sumar_ingresos_bomberos_gobernacion, sumar_ingresos_bomberos_sami


class PagosBomberoReport:
    def __init__(self, datos: pd.DataFrame, municipio, titulo_reporte, fecha_ini, fecha_fin, tipo_cta):
        self.datos = datos
        self.municipio = municipio
        self.fecha_ini = fecha_ini
        self.fecha_fin = fecha_fin
        self.titulo = titulo_reporte
        self.tipo_cta = tipo_cta
        self.fila_num = 0

    def generar_excel(self, ruta_salida="ingresos_bomberos.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Ingresos_bomberos"

        # ===============================
        # ⭐ ESTILOS
        # ===============================
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        title_font = Font(size=16, bold=True)
        date_font = Font(size=12, italic=True)
        total_font = Font(size=11, italic=True, bold=True)

        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        align_center = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # ===============================
        # ⭐ TITULO Y FECHA
        # ===============================
        titulo = f"{self.municipio['NombreMuni']} - {self.municipio['NombreDepto']}"
        titulo_rpt = f"{self.titulo}"
        fecha = f"Fecha de elaboración: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        periodo = f"Periodo desde el {self.fecha_ini} al {self.fecha_fin}"

        ws.merge_cells("A1:H1")
        ws["A1"] = titulo
        ws["A1"].alignment = align_center
        ws["A1"].font = title_font

        ws.merge_cells("A2:H4")
        ws["A2"] = titulo_rpt
        ws["A2"].alignment = align_center
        ws["A2"].font = title_font

        ws.merge_cells("A5:H5")
        ws["A5"] = fecha
        ws["A5"].alignment = align_center
        ws["A5"].font = date_font

        ws.merge_cells("A6:H6")
        ws["A6"] = periodo
        ws["A6"].alignment = align_center
        ws["A6"].font = date_font

        self.fila_num = 9  # Aquí irán los encabezados
        ws.freeze_panes = "A10"
        # ===============================
        # ⭐ ENCABEZADOS
        # ===============================
        columnas = []
        columnas.append("No.")

        # ===============================
        # ⭐ LLENADO DE DATOS
        # ===============================

        if self.tipo_cta:
            datos = sumar_ingresos_bomberos_sami(self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            for index, item in datos.iterrows():
                no = index+1
                num = item["No. Recibo"]
                dni = item["DNI"].strip()
                nombre_completo = item["Nombre Completo"]

                self.fila_num += 1

                ws.append([
                    no,
                    num,
                    dni,
                    nombre_completo,
                    f'{item["SALDO"]:,.2f}',
                    f'{item["RECUPERACION DE SALDOS"]:,.2f}',
                    f'{item["RECARGOS"]:,.2f}',
                    f'{item["INTERESES"]:,.2f}',
                    f'{item["TOTAL"]:,.2f}',

                ])
        else:
            datos = sumar_ingresos_bomberos_gobernacion(self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            for index, item in datos.iterrows():
                no = index+1
                num = item["No. Recibo"]
                dni = item["DNI"].strip()
                nombre_completo = item["Nombre Completo"]

                self.fila_num += 1

                ws.append([
                    no,
                    num,
                    dni,
                    nombre_completo,
                    f'{item["SALDO"]:,.2f}',
                    f'{item["RECUPERACION DE SALDOS"]:,.2f}',
                    f'{item["TOTAL"]:,.2f}',

                ])

        ws.cell(row=self.fila_num-1, column=2).font = total_font
        ws.cell(row=self.fila_num-1, column=3).font = total_font
        ws.cell(row=self.fila_num-1, column=4).font = total_font
        ws.cell(row=self.fila_num-1, column=5).font = total_font
        ws.cell(row=self.fila_num-1, column=6).font = total_font
        ws.cell(row=self.fila_num-1, column=7).font = total_font

        # ===============================
        # ⭐ AJUSTE AUTOMÁTICO DE COLUMNAS
        # ===============================
        for col in ws.columns:
            max_length = 0
            column = col[0].column  # número de columna
            column_letter = get_column_letter(column)

            for cell in col[8:]:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = max_length + 2  # un pequeño margen
            ws.column_dimensions[column_letter].width = adjusted_width

        # ===============================
        # ⭐ GUARDAR
        # ===============================
        wb.save(ruta_salida)
        return ruta_salida
