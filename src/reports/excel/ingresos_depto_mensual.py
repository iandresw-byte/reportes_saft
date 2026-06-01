import pandas as pd
from openpyxl.utils import get_column_letter
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from src.reports.tablas.ingresos_detallados_justicia_municipal import tabla_ingreso_diario_admin_excel_justicia
from src.reports.tablas.ingresos_detallados_procamut import tabla_ingreso_diario_admin_excel_procamut
from src.reports.tablas.ingresos_detallados_secretaria import tabla_ingreso_diario_admin_excel_secretaria
from src.reports.tablas.ingresos_detallados_uma import tabla_ingreso_diario_admin_excel_uma
from src.reports.tablas.ingresosd_detallados_desarrollo_urbano import tabla_ingreso_diario_admin_excel_urbanismo
from src.reports.utils.pintar_encabezado_excel import pintar_encabezado
from src.reports.tablas.ingresos_detallados_admin_triburaria_ip import tabla_ingreso_diario_admin_excel_tributaria_ip
from src.reports.tablas.ingresos_detallados_admin_triburaria_os import tabla_ingreso_diario_admin_excel_tributaria_os
from src.reports.tablas.ingresos_detallados_admin_tributaria_ics import tabla_ingreso_diario_admin_excel_tributaria_ics
from src.reports.tablas.ingresos_detallados_admin_triburaria_bi import tabla_ingreso_diario_admin_excel_tributaria_bi
from src.reports.tablas.ingresos_detallados_catastro import tabla_ingreso_diario_admin_excel_catastro
from src.reports.utils.manejador_data_depto_diario_frame import (sumar_ingresos_depto_diarios_catastro, sumar_ingresos_depto_diarios_admin_tributaria_bi,
                                                                 sumar_ingresos_depto_diarios_admin_tributaria_ics,
                                                                 sumar_ingresos_depto_diarios_admin_tributaria_ip,
                                                                 sumar_ingresos_depto_diarios_admin_tributaria_os,
                                                                 sumar_ingresos_depto_diarios_justicia,
                                                                 sumar_ingresos_depto_diarios_procamut,
                                                                 sumar_ingresos_depto_diarios_secretaria,
                                                                 sumar_ingresos_depto_diarios_uma,
                                                                 sumar_ingresos_depto_diarios_urbanizacion)


class IngresosDeptosMensualReport:
    def __init__(self, datos: pd.DataFrame, municipio, titulo_reporte, fecha_ini, fecha_fin, depto, tipo_impuesto):
        self.datos = datos
        self.municipio = municipio
        self.fecha_ini = fecha_ini
        self.fecha_fin = fecha_fin
        self.titulo = titulo_reporte
        self.fila_num = 0
        self.tipo_impuesto = tipo_impuesto
        self.departamento = depto

    def generar_excel(self, ruta_salida="ingresos_deptos_diario.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Ingresos_Depto_Det"

        # ===============================
        # ⭐ ESTILOS
        # ===============================

        title_font = Font(size=16, bold=True)
        date_font = Font(size=12, italic=True)
        total_font = Font(size=11, italic=True, bold=True)
        align_center = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

        # ===============================
        # ⭐ TITULO Y FECHA
        # ===============================
        titulo = f"{self.municipio['NombreMuni']} - {self.municipio['NombreDepto']}"
        titulo_rpt = f"{self.titulo}"
        fecha = f"Fecha de elaboración: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        periodo = f"Periodo desde el {self.fecha_ini} al {self.fecha_fin}"

        if self.tipo_impuesto == '0':
            letra_col = 'K'
        elif self.tipo_impuesto == '1':
            letra_col = 'M'
        elif self.tipo_impuesto == '2,3':
            letra_col = 'N'
        elif self.tipo_impuesto == '4':
            letra_col = 'M'
        else:
            letra_col = 'O'

        ws.merge_cells(f"A1:{letra_col}1")
        ws["A1"] = titulo
        ws["A1"].alignment = align_center
        ws["A1"].font = title_font

        ws.merge_cells(f"A2:{letra_col}4")
        ws["A2"] = titulo_rpt
        ws["A2"].alignment = align_center
        ws["A2"].font = title_font

        ws.merge_cells(f"A5:{letra_col}5")
        ws["A5"] = fecha
        ws["A5"].alignment = align_center
        ws["A5"].font = date_font

        ws.merge_cells(f"A6:{letra_col}6")
        ws["A6"] = periodo
        ws["A6"].alignment = align_center
        ws["A6"].font = date_font

        self.fila_num = 9  # Aquí irán los encabezados
        ws.freeze_panes = "A10"

        # ===============================
        # ⭐ LLENADO DE DATOS
        # ===============================
        columnas = []
        columnas.append("No.")
        if self.departamento == '1':
            datos = sumar_ingresos_depto_diarios_catastro(
                self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            ws = tabla_ingreso_diario_admin_excel_catastro(
                self, datos, ws)
        elif self.departamento == '2':
            if self.tipo_impuesto == '0':
                datos = sumar_ingresos_depto_diarios_admin_tributaria_os(
                    self.datos)
                ws = pintar_encabezado(self, datos, columnas, ws)
                self.fila_num += 1
                ws = tabla_ingreso_diario_admin_excel_tributaria_os(
                    self, datos, ws)
            elif self.tipo_impuesto == '1':
                datos = sumar_ingresos_depto_diarios_admin_tributaria_bi(
                    self.datos)
                ws = pintar_encabezado(self, datos, columnas, ws)
                self.fila_num += 1
                ws = tabla_ingreso_diario_admin_excel_tributaria_bi(
                    self, datos, ws)
            elif self.tipo_impuesto == '2,3':
                datos = sumar_ingresos_depto_diarios_admin_tributaria_ics(
                    self.datos)
                ws = pintar_encabezado(self, datos, columnas, ws)
                self.fila_num += 1
                ws = tabla_ingreso_diario_admin_excel_tributaria_ics(
                    self, datos, ws)
            elif self.tipo_impuesto == '4':
                datos = sumar_ingresos_depto_diarios_admin_tributaria_ip(
                    self.datos)
                ws = pintar_encabezado(self, datos, columnas, ws)
                self.fila_num += 1
                ws = tabla_ingreso_diario_admin_excel_tributaria_ip(
                    self, datos, ws)
            else:
                return None
        elif self.departamento == '3':
            datos = sumar_ingresos_depto_diarios_uma(
                self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            ws = tabla_ingreso_diario_admin_excel_uma(
                self, datos, ws)
        elif self.departamento == '4':
            datos = sumar_ingresos_depto_diarios_justicia(
                self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            ws = tabla_ingreso_diario_admin_excel_justicia(
                self, datos, ws)
        elif self.departamento == '6':
            datos = sumar_ingresos_depto_diarios_secretaria(
                self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            ws = tabla_ingreso_diario_admin_excel_secretaria(
                self, datos, ws)
        elif self.departamento == '7':
            datos = sumar_ingresos_depto_diarios_procamut(
                self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            ws = tabla_ingreso_diario_admin_excel_procamut(
                self, datos, ws)
        elif self.departamento == '8':
            datos = sumar_ingresos_depto_diarios_urbanizacion(
                self.datos)
            ws = pintar_encabezado(self, datos, columnas, ws)
            self.fila_num += 1
            ws = tabla_ingreso_diario_admin_excel_urbanismo(
                self, datos, ws)
        else:
            return None

        ws.cell(row=self.fila_num, column=2).font = total_font
        ws.cell(row=self.fila_num, column=3).font = total_font
        ws.cell(row=self.fila_num, column=4).font = total_font
        ws.cell(row=self.fila_num, column=5).font = total_font

        # ===============================
        # ⭐ AJUSTE AUTOMÁTICO DE COLUMNAS
        # ===============================
        for col in ws.columns:
            max_length = 0
            column = col[0].column  # número de columna
            column_letter = get_column_letter(column)

            for cell in col[8:]:
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                except:
                    pass
            adjusted_width = max_length + 2  # un pequeño margen
            ws.column_dimensions[column_letter].width = adjusted_width

        # ===============================
        # ⭐ GUARDAR
        # ===============================
        wb.save(ruta_salida)
        return ruta_salida
