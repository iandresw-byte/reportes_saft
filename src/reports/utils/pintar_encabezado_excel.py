from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def pintar_encabezado(excel, datos, columnas, ws):
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    align_center = Alignment(
        horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    for col in datos.columns:
        columnas.append(col)
        col_value = 0
        for col in range(1, len(columnas) + 1):
            c = ws.cell(row=excel.fila_num, column=col,
                        value=columnas[col_value])
            c.fill = header_fill
            c.font = header_font
            c.alignment = align_center
            c.border = border
            col_value += 1
    return ws
